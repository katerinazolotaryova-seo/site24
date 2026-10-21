#!/usr/bin/env python3
"""Generate a УКР invoice (рахунок-фактура) .xlsx from a requisites profile's
template (see LAYOUTS / requisites.json — each supplier profile has its own
template file and cell layout).

Usage example:
    python3 generate_invoice.py \
        --site zvlegal.com \
        --amount-uah 25339 \
        --date 2026-10-20 --suffix 1 \
        --requisites turchykhin \
        --out-dir .

If --service is omitted, it defaults to:
    "Консультування з питань інформатизації: SEO просування сайту <site>"
Pass --service explicitly to override (e.g. for a GEO consultation instead
of SEO promotion).

Either --amount-usd (converted via PrivatBank commercial rate) or
--amount-uah (a fixed UAH amount) must be given.

Output filename follows: рахунок_<service_code>_<site>_<MM-YYYY>.xlsx
"""
import argparse
import datetime
import json
import shutil
import subprocess
import sys
import tempfile
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

SKILL_DIR = Path(__file__).resolve().parent
REQUISITES_PATH = SKILL_DIR / "requisites.json"
SHEET_NAME = "Счет "  # NB: trailing space, matches the original templates

PRIVATBANK_COMMERCIAL_URL = "https://api.privatbank.ua/p24api/pubinfo?json&exchange&coursid=5"

# Each requisites profile (requisites.json) names a "layout" key below and a
# "template" file in this directory. A layout is a different physical invoice
# (different supplier => different cell positions/fields), not just different
# values plugged into the same cells — so add a new layout here whenever a
# new supplier's invoice has a different row/column structure, rather than
# trying to force it into an existing layout's cell map.
LAYOUTS = {
    "turchykhin_v1": {
        # requisite_cells: ordered (requisites.json key, cell) pairs for the
        # top "Постачальник" block plus the "Керівник" signature cell.
        "requisite_cells": [
            ("label", "C1"), ("edrpou", "C2"), ("iban", "C3"),
            ("bank", "C4"), ("bank_mfo", "C5"), ("director", "C20"),
        ],
        "invoice_no_cell": "A8",
        "invoice_no_fmt": "РАХУНОК-ФАКТУРА №{}",
        "payer_cell": "C10",
        "item_num_cell": "A15", "item_desc_cell": "B15", "item_unit_cell": "C15",
        "item_qty_cell": "D15", "item_price_cell": "E15", "item_sum_cell": "F15",
        "total_label_cell": "D16", "total_label_value": "Всього ",
        "total_sum_cell": "F16",
        "sum_label_cell": "B18", "sum_label_value": "Сума без ПДВ",
        "sum_value_cell": "D18", "sum_words_cell": "E18",
        "date_cell": "B23",
    },
    "zolotaryova_v1": {
        "requisite_cells": [
            ("label", "C1"), ("iban", "C2"), ("bank", "C3"), ("mfo", "C4"),
            ("inn", "C5"), ("address", "C6"), ("phone", "C7"), ("director", "C23"),
        ],
        "invoice_no_cell": "A11",
        "invoice_no_fmt": "РАХУНОК-ФАКТУРА №{}",
        "payer_cell": "C13",
        "item_num_cell": "A18", "item_desc_cell": "B18", "item_unit_cell": "C18",
        "item_qty_cell": "D18", "item_price_cell": "E18", "item_sum_cell": "F18",
        "total_label_cell": "D19", "total_label_value": "Всього ",
        "total_sum_cell": "F19",
        "sum_label_cell": "B21", "sum_label_value": "Сума без ПДВ",
        "sum_value_cell": "D21", "sum_words_cell": "E21",
        "date_cell": "B26",
    },
}


# ---- Ukrainian number-to-words (for the "сума словами" line) ----

_ONES_M = ["", "один", "два", "три", "чотири", "п'ять", "шість", "сім", "вісім", "дев'ять"]
_ONES_F = ["", "одна", "дві", "три", "чотири", "п'ять", "шість", "сім", "вісім", "дев'ять"]
_TEENS = ["десять", "одинадцять", "дванадцять", "тринадцять", "чотирнадцять",
          "п'ятнадцять", "шістнадцять", "сімнадцять", "вісімнадцять", "дев'ятнадцять"]
_TENS = ["", "", "двадцять", "тридцять", "сорок", "п'ятдесят",
         "шістдесят", "сімдесят", "вісімдесят", "дев'яносто"]
_HUNDREDS = ["", "сто", "двісті", "триста", "чотириста", "п'ятсот",
             "шістсот", "сімсот", "вісімсот", "дев'ятсот"]
_SCALES = [
    ("", "", ""),
    ("тисяча", "тисячі", "тисяч"),
    ("мільйон", "мільйони", "мільйонів"),
    ("мільярд", "мільярди", "мільярдів"),
]


def _scale_form(n, forms):
    if 11 <= n % 100 <= 14:
        return forms[2]
    last = n % 10
    if last == 1:
        return forms[0]
    if 2 <= last <= 4:
        return forms[1]
    return forms[2]


def _group_words(n, feminine):
    words = []
    h, r = divmod(n, 100)
    if h:
        words.append(_HUNDREDS[h])
    if 10 <= r < 20:
        words.append(_TEENS[r - 10])
    else:
        t, o = divmod(r, 10)
        if t:
            words.append(_TENS[t])
        if o:
            words.append((_ONES_F if feminine else _ONES_M)[o])
    return words


def number_to_words_uk(n: int, base_feminine: bool = False) -> str:
    """base_feminine: grammatical gender of the noun counted by the final
    (units) group — e.g. True for "гривня" ("одна гривня", not "один гривня").
    The thousands group is always feminine ("тисяча"); millions/billions
    are always masculine, regardless of base_feminine."""
    if n == 0:
        return "нуль"
    groups = []
    rem = n
    while rem > 0:
        groups.append(rem % 1000)
        rem //= 1000
    parts = []
    for i in reversed(range(len(groups))):
        g = groups[i]
        if g == 0:
            continue
        if i == 0:
            feminine = base_feminine
        elif i == 1:
            feminine = True  # "тисяча"
        else:
            feminine = False  # "мільйон", "мільярд"
        words = _group_words(g, feminine)
        if i > 0:
            words.append(_scale_form(g, _SCALES[i]))
        parts.extend(words)
    return " ".join(parts)


def amount_in_words_uah(amount: float) -> str:
    grn = int(amount)
    kop = round((amount - grn) * 100)
    if kop == 100:
        grn, kop = grn + 1, 0
    grn_words = number_to_words_uk(grn, base_feminine=True).capitalize() if grn else "Нуль"
    grn_form = _scale_form(grn, ("гривня", "гривні", "гривень"))
    kop_form = _scale_form(kop, ("копійка", "копійки", "копійок"))
    return f"грн ({grn_words} {grn_form} {kop:02d} {kop_form})"


# ---- PrivatBank exchange rate ----

def fetch_usd_uah_rate(rate_type: str) -> float:
    with urllib.request.urlopen(PRIVATBANK_COMMERCIAL_URL, timeout=10) as resp:
        data = json.load(resp)
    for row in data:
        if row.get("ccy") == "USD":
            key = "sale" if rate_type == "sale" else "buy"
            return float(row[key])
    raise RuntimeError("USD rate not found in PrivatBank response")


# ---- Recalculate formulas via headless LibreOffice ----

def recalc_formulas(xlsx_path: Path):
    # soffice refuses to write back over the same path it read from, so
    # convert into a separate temp output dir, then copy the recalculated
    # bytes over the real (final) filename.
    with tempfile.TemporaryDirectory() as src_dir, tempfile.TemporaryDirectory() as out_dir:
        tmp_in = Path(src_dir) / "invoice.xlsx"
        shutil.copy(xlsx_path, tmp_in)
        subprocess.run(
            ["soffice", "--headless", "--norestore", "--convert-to", "xlsx",
             "--outdir", out_dir, str(tmp_in)],
            check=True, capture_output=True, timeout=60,
        )
        shutil.copy(Path(out_dir) / "invoice.xlsx", xlsx_path)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--site", required=True, help="e.g. frigat.net")
    p.add_argument("--service", default=None,
                    help="Full service text written into the invoice row. "
                         "Defaults to \"Консультування з питань інформатизації: "
                         "SEO просування сайту <site>\" if omitted.")
    p.add_argument("--unit", default="місяць")
    p.add_argument("--service-code", default=None,
                    help="Short label for the filename (e.g. 'SEO', 'GEO'). "
                         "Defaults to 'SEO'/'GEO' auto-detected from --service, "
                         "else a slug of --service.")
    p.add_argument("--qty", type=float, default=1)
    p.add_argument("--amount-usd", type=float)
    p.add_argument("--amount-uah", type=float)
    p.add_argument("--rate", choices=["sale", "buy"], default="sale",
                    help="PrivatBank commercial rate side to use for USD conversion (default: sale)")
    p.add_argument("--date", default=datetime.date.today().isoformat(),
                    help="YYYY-MM-DD, defaults to today")
    p.add_argument("--suffix", default="1", help="Invoice number suffix after '/'")
    p.add_argument("--requisites", default="turchykhin")
    p.add_argument("--payer", default="")
    p.add_argument("--out-dir", default=".")
    args = p.parse_args()

    if (args.amount_usd is None) == (args.amount_uah is None):
        sys.exit("Specify exactly one of --amount-usd or --amount-uah")

    if args.service is None:
        args.service = f"Консультування з питань інформатизації: SEO просування сайту {args.site}"

    invoice_date = datetime.date.fromisoformat(args.date)

    requisites = json.loads(REQUISITES_PATH.read_text())
    if args.requisites not in requisites:
        sys.exit(f"Unknown requisites profile '{args.requisites}'. "
                  f"Known: {', '.join(requisites)}")
    req = requisites[args.requisites]
    layout = LAYOUTS[req["layout"]]
    template_path = SKILL_DIR / req["template"]

    if args.amount_usd is not None:
        rate = fetch_usd_uah_rate(args.rate)
        amount_uah = round(args.amount_usd * rate, 2)
    else:
        amount_uah = round(args.amount_uah, 2)

    invoice_no = f"{invoice_date.strftime('%d%m%y')}/{args.suffix}"

    wb = load_workbook(template_path)
    ws = wb[SHEET_NAME]

    for key, cell in layout["requisite_cells"]:
        ws[cell] = req[key]

    ws[layout["invoice_no_cell"]] = layout["invoice_no_fmt"].format(invoice_no)
    ws[layout["payer_cell"]] = args.payer or None

    ws[layout["item_num_cell"]] = 1
    ws[layout["item_desc_cell"]] = args.service
    ws[layout["item_unit_cell"]] = args.unit
    ws[layout["item_qty_cell"]] = args.qty
    ws[layout["item_price_cell"]] = amount_uah
    ws[layout["item_sum_cell"]] = f"={layout['item_price_cell']}"
    ws[layout["total_label_cell"]] = layout["total_label_value"]
    ws[layout["total_sum_cell"]] = f"={layout['item_sum_cell']}"
    ws[layout["sum_label_cell"]] = layout["sum_label_value"]
    ws[layout["sum_value_cell"]] = f"={layout['total_sum_cell']}"
    ws[layout["sum_words_cell"]] = amount_in_words_uah(amount_uah)
    ws[layout["date_cell"]] = datetime.datetime.combine(invoice_date, datetime.time())

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    service_code = args.service_code
    if service_code is None:
        for code in ("SEO", "GEO"):
            if code.lower() in args.service.lower():
                service_code = code
                break
    if service_code is None:
        # Fallback: strip the site name out of the service text, wherever it
        # appears ("для <site>", "сайту <site>", ...), and tidy up leftover
        # punctuation/whitespace for a filesystem-safe label.
        service_code = args.service.replace(args.site, "").replace(":", "").strip()
        service_code = " ".join(service_code.split())
    filename = f"рахунок_{service_code}_{args.site}_{invoice_date.strftime('%m-%Y')}.xlsx"
    out_path = out_dir / filename
    wb.save(out_path)

    recalc_formulas(out_path)

    print(out_path)


if __name__ == "__main__":
    main()
