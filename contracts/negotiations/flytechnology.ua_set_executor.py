"""Міняє Виконавця на ФОП Золотарьову та дату на сьогодні у фінальному договорі flytechnology.ua."""
import copy, datetime as dt, sys
from pathlib import Path
import docx, yaml

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from build import set_date_runs, find_par  # noqa: E402

src, out = sys.argv[1], sys.argv[2]
ent = yaml.safe_load((Path(__file__).resolve().parents[1] / "entities/zolotaryova.yaml").read_text(encoding="utf-8"))
doc = docx.Document(src)
today = dt.date.today()

set_date_runs(find_par(doc, "м. Київ").runs, today)
set_date_runs(find_par(doc, "до Договору", start=True).runs, today)
# жовте виділення дат у шаблоні — позначка «оновити дату»; після оновлення прибираємо
for p in (find_par(doc, "м. Київ"), find_par(doc, "до Договору", start=True)):
    for r in p.runs:
        r.font.highlight_color = None

pre = find_par(doc, "«Виконавець»")
pre.runs[0].text = ent["name"]
pre.runs[1].text = " ".join(ent["preamble_tail"].split())

for table in doc.tables:
    row = next(r for r in table.rows if r.cells[0].text.strip().startswith("ФОП"))
    pars = row.cells[0].paragraphs
    sig = next(k for k, p in enumerate(pars) if p.text.startswith("____"))
    lines = ent["requisites"]
    assert len(lines) <= sig
    # рядки 1..N — у форматі звичайного рядка (pars[2]), рядок 0 — жирна назва
    for k in range(1, len(lines)):
        if k != 2:
            pars[k]._p.addprevious(copy.deepcopy(pars[2]._p))
            pars[k]._p.getparent().remove(pars[k]._p)
    pars = row.cells[0].paragraphs
    for k in range(sig):
        p = pars[k]
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        text = lines[k] if k < len(lines) else ""
        if p.runs:
            p.runs[0].text = text
        elif text:
            p.add_run(text)
    pars[sig].runs[-1].text = ent["signature"]

doc.save(out)
