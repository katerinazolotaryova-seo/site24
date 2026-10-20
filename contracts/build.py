#!/usr/bin/env python3
"""Збирає договір з шаблону templates/<service>.docx і файлу замовлення orders/*.yaml.

    python3 contracts/build.py contracts/orders/2026-10-20_zvlegal.com.yaml

Результат: contracts/output/Договір_<SERVICE>_<site>_<MM-YYYY>.docx
"""
import copy
import datetime as dt
import sys
from pathlib import Path

import docx
import yaml

ROOT = Path(__file__).resolve().parent

MONTHS = ["січня", "лютого", "березня", "квітня", "травня", "червня", "липня",
          "серпня", "вересня", "жовтня", "листопада", "грудня"]

ONES_M = ["", "один", "два", "три", "чотири", "п’ять", "шість", "сім", "вісім", "дев’ять"]
ONES_F = ["", "одна", "дві", "три", "чотири", "п’ять", "шість", "сім", "вісім", "дев’ять"]
TEENS = ["десять", "одинадцять", "дванадцять", "тринадцять", "чотирнадцять", "п’ятнадцять",
         "шістнадцять", "сімнадцять", "вісімнадцять", "дев’ятнадцять"]
TENS = ["", "", "двадцять", "тридцять", "сорок", "п’ятдесят", "шістдесят", "сімдесят",
        "вісімдесят", "дев’яносто"]
HUNDREDS = ["", "сто", "двісті", "триста", "чотириста", "п’ятсот", "шістсот", "сімсот",
            "вісімсот", "дев’ятсот"]


def plural(n, one, few, many):
    if n % 100 in range(11, 20):
        return many
    return {1: one, 2: few, 3: few, 4: few}.get(n % 10, many)


def triple(n, ones):
    words = [HUNDREDS[n // 100]]
    rest = n % 100
    if 10 <= rest < 20:
        words.append(TEENS[rest - 10])
    else:
        words += [TENS[rest // 10], ones[rest % 10]]
    return [w for w in words if w]


def words_uk(n):
    """Ціле число прописом (до 999 999 999)."""
    if n == 0:
        return "нуль"
    parts = []
    for div, ones, forms in ((10**6, ONES_M, ("мільйон", "мільйони", "мільйонів")),
                             (10**3, ONES_F, ("тисяча", "тисячі", "тисяч"))):
        chunk = n // div % 1000
        if chunk:
            parts += triple(chunk, ones) + [plural(chunk, *forms)]
    parts += triple(n % 1000, ONES_M)
    return " ".join(parts)


def fmt_amount(n):
    return f"{n:,}".replace(",", " ")


def price_text(price):
    amount = int(price["amount"])
    spelled = words_uk(amount)
    if price["currency"].upper() == "USD":
        return (f"суму в гривнях, еквівалентну {fmt_amount(amount)} ({spelled}) "
                f"{plural(amount, 'долару', 'доларам', 'доларам')} США за комерційним курсом "
                f"продажу АТ КБ «ПриватБанк» на дату виставлення рахунку, звіт за якими "
                f"надається після виконання Додатку 1 та/або за запитом Замовника.")
    return (f"{fmt_amount(amount)} ({spelled}) {plural(amount, 'гривню', 'гривні', 'гривень')} "
            f"на дату виставлення рахунку, звіт за якими надається після виконання "
            f"Додатку 1 та/або за запитом Замовника.")


def find_par(doc, needle, start=False):
    for p in doc.paragraphs:
        if (p.text.startswith(needle) if start else needle in p.text):
            return p
    raise SystemExit(f"У шаблоні не знайдено абзац з текстом: {needle!r}")


def set_date_runs(runs, date):
    """Run-и дати у шаблоні: '12', ..., 'жовтня ', '20', '26'."""
    texts = [r.text for r in runs]
    i = texts.index("«") + 1 if "«" in texts else next(
        k for k, t in enumerate(texts) if t.strip().isdigit())
    runs[i].text = f"{date.day:02d}"
    month = next(k for k in range(i, len(runs)) if runs[k].text.strip() in MONTHS)
    runs[month].text = MONTHS[date.month - 1] + " "
    runs[month + 1].text = str(date.year)[:2]
    runs[month + 2].text = str(date.year)[2:]


def fill_party_cell(exec_cell, client_cell, lines, signature):
    """Заповнює колонку ЗАМОВНИК, копіюючи форматування рядків колонки ВИКОНАВЕЦЬ."""
    exec_pars = exec_cell.paragraphs
    sig_idx = next(k for k, p in enumerate(exec_pars) if p.text.startswith("____"))
    if len(lines) > sig_idx - 1:
        raise SystemExit(f"Забагато рядків реквізитів клієнта (максимум {sig_idx - 1}).")
    content = lines + [""] * (sig_idx - len(lines))
    for p in list(client_cell.paragraphs):
        p._p.getparent().remove(p._p)
    for k, text in enumerate(content):
        src = exec_pars[min(k, len(exec_pars) - 1)]
        new = copy.deepcopy(src._p)
        client_cell._tc.append(new)
        par = client_cell.paragraphs[-1]
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)
        if par.runs:
            par.runs[0].text = text
        elif text:
            par.add_run(text)
    # Рядок підпису: «____/» + прізвище у тому ж форматуванні, що й у Виконавця
    sig = copy.deepcopy(exec_pars[sig_idx]._p)
    client_cell._tc.append(sig)
    client_cell.paragraphs[-1].runs[-1].text = signature


def build(order_path):
    order = yaml.safe_load(Path(order_path).read_text(encoding="utf-8"))
    service = order["service"].lower()
    date = order.get("date") or dt.date.today()
    if isinstance(date, str):
        date = dt.date.fromisoformat(date)
    number = order.get("number") or f"{date:%d%m%y}/1"
    entity = yaml.safe_load((ROOT / "entities" / f"{order['entity']}.yaml").read_text(encoding="utf-8"))
    if entity["id"] != "turchykhin":
        raise SystemExit("Шаблон поки містить лише реквізити Турчихіна — "
                         "для іншої особи треба додати заміну блоку Виконавця.")
    client = order["client"]

    doc = docx.Document(ROOT / "templates" / f"{service}.docx")

    # Номер і дата договору
    head = find_par(doc, "Договір", start=True)
    head.runs[1].text = f" №{number}"
    set_date_runs(find_par(doc, "м. Київ").runs, date)

    # Преамбула: замість «___» — Замовник
    pre = find_par(doc, "іменований надалі")
    preamble = " ".join(client["preamble"].split())
    bold = client.get("bold_name", "")
    pre.runs[0].text = bold
    pre.runs[1].text = preamble[len(bold):] + pre.runs[1].text.replace("«3амовник»", "«Замовник»")

    # Сайт у п. 3.2.1
    site_par = find_par(doc, "бренду Замовника")
    r = site_par.runs[0]
    r.text = r.text[:r.text.index("бренду Замовника") + len("бренду Замовника")] + " " + order["site"]

    # Додаток 1: номер і дата
    app = find_par(doc, "до Договору", start=True)
    num_run = next(r for r in app.runs if "/" in r.text)
    num_run.text = number
    set_date_runs(app.runs, date)

    # Додаток 1: план робіт (нумерований список після заголовка)
    plan_head = find_par(doc, "План робіт")
    items = []
    el = plan_head._p.getnext()
    while el is not None and el.tag.endswith("}p"):
        par = docx.text.paragraph.Paragraph(el, plan_head._parent)
        if not par.text.strip() or par.text.startswith("ІІ."):
            break
        items.append(par)
        el = el.getnext()
    plan = order["plan"]
    while len(items) < len(plan):
        new = copy.deepcopy(items[-1]._p)
        items[-1]._p.addnext(new)
        items.append(docx.text.paragraph.Paragraph(new, plan_head._parent))
    for par, text in zip(items, plan):
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)
        par.runs[0].text = text
    for par in items[len(plan):]:
        par._p.getparent().remove(par._p)

    # Додаток 1: вартість
    cost = find_par(doc, "Вартість Послуг")
    cost.runs[0].text = "Вартість Послуг (робіт), що зазначені в цьому Додатку, складає"
    cost.runs[2].text = price_text(order["price"])

    # Таблиці реквізитів (договір і додаток)
    for table in doc.tables:
        row = next(r for r in table.rows if r.cells[0].text.strip().startswith("ФОП"))
        fill_party_cell(row.cells[0], row.cells[1], client["requisites"], client["signature"])
        for r in table.rows:  # прибрати зайві «___» у порожніх рядках під таблицею
            if r is not row and r.cells[1].text.strip() == "___":
                for p in r.cells[1].paragraphs:
                    for run in p.runs:
                        run.text = ""

    out_dir = ROOT / "output"
    out_dir.mkdir(exist_ok=True)
    out = out_dir / f"Договір_{service.upper()}_{order['site']}_{date:%m-%Y}.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    print(build(sys.argv[1]))
