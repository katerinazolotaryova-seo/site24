"""Договір SEO для novagruz.com на ФОП Лубіну (шаблон templates/seo_lubina.docx)."""
import copy, sys
from pathlib import Path
import docx

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from build import find_par  # noqa: E402

NUMBER, DAY, MONTH, YEAR = "301026/1", "30", "жовтня", "2026"
SITE = "novagruz.com"
CLIENT = "Михайлов Михайло Володимирович"
CLIENT_LINES = [
    "ФОП Михайлов Михайло Володимирович",
    "РНОКПП 3269719850",
    "IBAN UA963052990000026007040813510",
    "в АТ КБ «ПРИВАТБАНК»",
    "E-mail: Miheylux@gmail.com",
]
CLIENT_SIG = "Михайлов М.В."
PLAN = [
    "Розробка семантичного ядра під пріоритетні напрямки бізнесу — до 40 тематичних груп запитів, з перевіркою реального пошукового попиту (а не лише інтуїтивних формулювань) і без географічної плутанини, яка зараз розмиває сигнал релевантності сайту",
    "Вирішення питання про створення RU-версії сайту",
    "Аудит структури сайту та формування технічного завдання на її розширення — визначення, яких сторінок послуг і напрямків бракує порівняно з лідерами ніші, щоб сайт міг претендувати на трафік, який зараз повністю йде конкурентам",
    "Розробка мета-тегів на основі зібраної семантики — з чіткою геолокаційною сегментацією, що усуває поточну плутанину сигналів для пошукової системи",
    "Формування технічних завдань на написання й оптимізацію текстів для сторінок послуг — 10 текстів, підготовлених на основі семантичного ядра та детального аналізу контентної стратегії прямих конкурентів",
    "Повний технічний SEO-аудит сайту з формуванням технічного завдання для розробника — виявлення й усунення технічних бар'єрів, що стримують індексацію та ранжування",
    "Аналіз E-E-A-T факторів (Experience, Expertise, Authoritativeness, Trustworthiness) провідних гравців ніші та формування технічного завдання на підвищення рівня довіри до сайту в очах користувачів і пошукових систем",
]
# 580 USD × 45,8716 (комерційний курс продажу ПриватБанку, 30.10.2026) = 26 605,53 → 26 606 грн
PRICE = "Вартість Послуг (робіт), що зазначений в цьому Додатку складає 26\u00a0606 грн, звіт за якими надається після виконання Додатку 1 та/або за запитом Замовника."


def set_text(p, text):
    runs = [r for r in p.runs if r.text] or p.runs
    for r in p.runs:
        if r._r is not runs[0]._r:
            r._r.getparent().remove(r._r)
    runs[0].text = text


src, out = sys.argv[1], sys.argv[2]
doc = docx.Document(src)

# Номер і дата
find_par(doc, "Договір №", start=True).runs[1].text = NUMBER
date_run = next(r for r in find_par(doc, "м. Київ").runs if "«" in r.text)
date_run.text = f"   «{DAY}» {MONTH} {YEAR} р."

# Преамбула
pre = find_par(doc, "«Виконавець»")
pre.runs[1].text = CLIENT
pre.runs[3].text = pre.runs[3].text.replace("«3амовник»", "«Замовник»")

# Сайт: гіперпосилання в п. 2.1 + п. 3.2.1
p21 = find_par(doc, "2.1.", start=True)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
hl = p21._p.find(W + "hyperlink")
hl.find(f"{W}r/{W}t").text = SITE
doc.part.rels[hl.get(R + "id")]._target = f"http://{SITE}/"
p321 = find_par(doc, "3.2.1.", start=True)
p321.runs[0].text = p321.runs[0].text.rstrip() + " " + SITE + "."

# Строк дії
p112 = find_par(doc, "11.2.", start=True)
next(r for r in p112.runs if "31.12." in r.text).text = "до 31.12.2027"

# Додаток 1: номер, дата
app = find_par(doc, "до Договору №", start=True)
app.runs[1].text = NUMBER + " "
app.runs[3].text = DAY
app.runs[5].text = f" {MONTH} "
app.runs[6].text, app.runs[7].text = YEAR[:2], YEAR[2:]

# Додаток 1: план
head = find_par(doc, "І. План робіт", start=True)
items, el = [], head._p.getnext()
while el is not None:
    p = docx.text.paragraph.Paragraph(el, head._parent)
    if not p.text.strip():
        break
    items.append(p)
    el = el.getnext()
while len(items) < len(PLAN):
    new = copy.deepcopy(items[-1]._p)
    items[-1]._p.addnext(new)
    items.append(docx.text.paragraph.Paragraph(new, head._parent))
for p, t in zip(items, PLAN):
    set_text(p, t)

set_text(find_par(doc, "Вартість Послуг", start=True), PRICE)

# Реквізити Замовника в обох таблицях
for table in doc.tables:
    row = next(r for r in table.rows if r.cells[1].text.strip().startswith("ФОП"))
    pars = row.cells[1].paragraphs
    sig = next(k for k, p in enumerate(pars) if p.text.startswith("____"))
    assert len(CLIENT_LINES) < sig
    for k in range(sig):
        p = pars[k]
        if 0 < k < len(CLIENT_LINES) and k != 1:  # формат звичайного рядка
            new = copy.deepcopy(pars[1]._p)
            p._p.addprevious(new)
            p._p.getparent().remove(p._p)
    pars = row.cells[1].paragraphs
    for k in range(sig):
        text = CLIENT_LINES[k] if k < len(CLIENT_LINES) else ""
        if pars[k].runs:
            set_text(pars[k], text)
        elif text:
            pars[k].add_run(text)
    next(r for r in pars[sig].runs if r.text.strip() and not r.text.startswith("_")).text = CLIENT_SIG

doc.save(out)
